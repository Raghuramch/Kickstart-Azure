"""Export generated Azure VMs and all returned details to an Excel workbook."""

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


HERE = Path(__file__).resolve().parent


def safe_cell(value):
    """Keep Azure values as text and prevent spreadsheet formula execution."""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    elif value is None:
        value = ""
    else:
        value = str(value)
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def flatten(value, prefix=""):
    if isinstance(value, dict):
        for key, item in value.items():
            yield from flatten(item, f"{prefix}.{key}" if prefix else key)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from flatten(item, f"{prefix}[{index}]")
    else:
        yield prefix, value


def main():
    config = json.loads((HERE / "export_config.json").read_text(encoding="utf-8"))
    machines = []
    for name in config["vm_names"]:
        command = [
            "az", "vm", "show", "--name", name,
            "--resource-group", config["resource_group_name"],
            "--subscription", config["subscription_id"],
            "--show-details", "--output", "json",
        ]
        try:
            result = subprocess.run(command, check=True, capture_output=True, text=True)
        except FileNotFoundError:
            sys.exit("Azure CLI (az) is required. Install it and run az login first.")
        except subprocess.CalledProcessError as error:
            sys.exit(f"Could not read VM {name}: {error.stderr.strip()}")
        machines.append(json.loads(result.stdout))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "VMs"
    tag_sheet = workbook.create_sheet("Tags")
    detail_sheet = workbook.create_sheet("All Details")
    tag_keys = sorted({key for vm in machines for key in (vm.get("tags") or {})})
    summary.append([
        "Subscription ID", "Resource Group", "VM Name", "VM ID", "Location",
        "VM Size", "OS Type", "Private IP", "Public IP", "Power State",
        *[f"Tag: {key}" for key in tag_keys],
    ])
    tag_sheet.append(["VM Name", "Tag Key", "Tag Value"])
    detail_sheet.append(["VM Name", "Field", "Value"])

    for vm in machines:
        tags = vm.get("tags") or {}
        summary.append([safe_cell(value) for value in (
            config["subscription_id"], config["resource_group_name"], vm.get("name"),
            vm.get("id"), vm.get("location"),
            (vm.get("hardwareProfile") or {}).get("vmSize"),
            ((vm.get("storageProfile") or {}).get("osDisk") or {}).get("osType"),
            vm.get("privateIps"), vm.get("publicIps"), vm.get("powerState"),
            *[tags.get(key) for key in tag_keys],
        )])
        for key, value in sorted(tags.items()):
            tag_sheet.append([safe_cell(vm.get("name")), safe_cell(key), safe_cell(value)])
        for field, value in flatten(vm):
            detail_sheet.append([safe_cell(vm.get("name")), safe_cell(field), safe_cell(value)])

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(60, max(14, max(len(str(cell.value or "")) for cell in column) + 2))

    output = HERE / "azure_vm_tags.xlsx"
    workbook.save(output)
    print(f"Exported {len(machines)} VM(s) to {output}")


if __name__ == "__main__":
    main()
